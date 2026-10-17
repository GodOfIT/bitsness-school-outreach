# -*- coding: utf-8 -*-
import sys
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
from data import all_rows, CRIT
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.comments import Comment

OUT = os.path.join(ROOT, 'output', 'Bang_uu_tien_tiep_can_truong.xlsx')
os.makedirs(os.path.dirname(OUT), exist_ok=True)
F = "Arial"
TEAL = "1F4E5F"
INPUT = PatternFill("solid", fgColor="FFF2CC")
KEYIN = PatternFill("solid", fgColor="FFD966")
HEAD = PatternFill("solid", fgColor=TEAL)
SUB = PatternFill("solid", fgColor="DCE8EC")
CALC = PatternFill("solid", fgColor="EEF3F5")
thin = Side(style="thin", color="C9D3D8")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=F, bold=bold, color=color, size=size, italic=italic)


def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font = f(True, "FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f(True, TEAL, 15)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(False, "555555", 10, True)


rows = all_rows()
rows.sort(key=lambda r: -r['score'])
N = len(rows)
FIRST = 5
LAST = FIRST + N - 1

wb = Workbook()

# ======================= 1. ĐỌC TRƯỚC =======================
ws0 = wb.active
ws0.title = "Đọc trước"
ws0.column_dimensions["A"].width = 4
ws0.column_dimensions["B"].width = 120
title(ws0, "Bảng ưu tiên tiếp cận trường – Bitsness (LMS, số hóa học liệu, khóa ngoại ngữ/khoa học/mô phỏng)",
      "Lập ngày 17/10/2026. Điểm hiện tại là SƠ BỘ, cần cập nhật sau khi bạn điền quan hệ (C8) và xác minh dữ liệu.")
blocks = [
    ("MỤC ĐÍCH", [
        "Xếp hạng 97 cơ sở đào tạo ở Hà Nội (51 cao đẳng, 30 đại học, 2 trung cấp, 14 TT GDNN-GDTX) để quyết định trường nào liên hệ trước.",
        "Mỗi trường được chấm 11 tiêu chí theo thang 1–5, nhân trọng số, quy về thang 100 và tự xếp vào Đợt 1/2/3/4.",
    ]),
    ("CÁCH DÙNG (4 BƯỚC)", [
        "1. Sheet 'Trọng số': xem lại trọng số 11 tiêu chí (ô vàng) và ngưỡng chia đợt. Tổng trọng số nên = 100.",
        "2. Sheet 'Chấm điểm': điền cột C8 'Đường tiếp cận' (ô vàng đậm) cho từng trường. Mặc định = 1 (chưa quen ai). Đây là tiêu chí nặng nhất (12%), điền xong thứ hạng sẽ thay đổi đáng kể.",
        "3. Trước khi gọi một trường, xác minh các dòng có cờ đỏ/cam ở cột 'Độ tin cậy dữ liệu', rồi sửa điểm C1–C11 cho đúng thực tế.",
        "4. Sheet 'Pipeline': đã điền sẵn các trường Đợt 1 và Đợt 2. Cập nhật giai đoạn sau mỗi lần liên hệ; sheet 'Kế hoạch' tự đếm tiến độ phễu.",
        "Sau mỗi cuộc gặp, dùng sheet 'Câu hỏi khảo sát' để hỏi đúng các ý cần cho 11 tiêu chí, rồi quay lại sửa điểm.",
    ]),
    ("CHÚ GIẢI", [
        "Ô nền vàng nhạt, chữ xanh dương: dữ liệu bạn nhập/sửa.  Ô vàng đậm (cột C8): ưu tiên điền.  Ô nền xám xanh: công thức, không gõ đè.",
        "Cột 'Độ tin cậy dữ liệu': đỏ = sai/không xác minh được; cam = chưa xác minh hoặc ước lượng; xanh = đã xác minh qua nguồn chính thức.",
        "Tổng điểm = SUMPRODUCT(điểm C1..C11 × trọng số) / (5 × tổng trọng số) × 100.  Đợt = theo ngưỡng ở sheet 'Trọng số'; trường ngoài Hà Nội luôn vào Đợt 4.",
        "Cột 'Gói chào mở cửa' tự đề xuất theo C1 (khoảng trống LMS), C2 (mô phỏng) và C3 (ngoại ngữ).",
    ]),
    ("BỐI CẢNH CHÍNH SÁCH ẢNH HƯỞNG ĐẾN THỨ TỰ ƯU TIÊN", [
        "• Kế hoạch 292/KH-UBND (27/7/2026): Hà Nội gộp 10 cao đẳng + 9 trung cấp công lập thành 8 cao đẳng + 2 trung cấp; 29 TT GDNN-GDTX còn 14; hoàn tất trước 25/8/2026. Trường mới gộp phải hợp nhất chương trình, học liệu và dữ liệu sinh viên của 2–3 trường cũ, đúng với LMS và số hóa học liệu. Vì vậy các trường này được C7 = 5.",
        "   Nguồn: giaoduc.net.vn/ha-noi-hoan-tat-sap-xep-co-so-gdnn-gdtx-dai-hoc-cong-lap-truoc-ngay-2582026-post261701.gd ; vov.vn/xa-hoi/ha-noi-sap-xep-to-chuc-lai-cac-co-so-giao-duc-nghe-nghiep-giao-duc-thuong-xuyen-post1319056.vov",
        "• Thông báo 750-TB/TU (đưa tin 6/10/2026): bố trí BGH mới cho ĐH Thủ đô và 8 trường cao đẳng trực thuộc TP. BGH mới thường muốn có kết quả nhanh và chưa bị ràng buộc với nhà cung cấp cũ. Tên Hiệu trưởng xem cột 'Ghi chú' trong sheet Chấm điểm.",
        "   Nguồn: giaoduc.net.vn/ha-noi-bo-tri-nhan-su-lanh-dao-cac-truong-dai-hoc-cao-dang-truc-thuoc-post263057.gd",
        "• QĐ 2371/QĐ-TTg (27/10/2025), đề án đưa tiếng Anh thành ngôn ngữ thứ hai giai đoạn 2025–2035, áp dụng cho cả GDNN và đại học. Mục tiêu đến 2030: ≥15% cơ sở GDNN và ≥20% cơ sở đại học đạt mức 1. Đề án có nhiệm vụ về học liệu số mở, nền tảng học tập số dùng AI và khuyến khích hợp tác với doanh nghiệp giáo dục số. Đây là căn cứ cho C3 và cho thông điệp chào khóa ngoại ngữ.",
        "   Nguồn: luatvietnam.vn/giao-duc/quyet-dinh-2371-qd-ttg-2025-phe-duyet-de-an-tieng-anh-thanh-ngon-ngu-thu-hai-trong-truong-hoc-416990-d1.html",
    ]),
    ("THAY ĐỔI SO VỚI FILE 'Các trường cao đẳng' CỦA BẠN", [
        "• Gộp #9 và #16 (CĐ Công nghệ cao Hà Nội, 2 cơ sở của cùng một trường) thành 1 dòng.",
        "• #1 CĐ Y tế Hà Đông đã hợp nhất vào CĐ Y tế Hà Nội, đổi tên dòng. #15 CĐ Cộng đồng Hà Nội đã sáp nhập vào CĐ Thương mại và Du lịch HN (#33), gộp vào dòng đó.",
        "• Chuyển sang 'Ngoài HN' (Đợt 4): #10 iSPACE (đặt tại TP.HCM), #44 CĐ PT-TH I (Phủ Lý), #52 CĐ Cơ khí Nông nghiệp (Bình Xuyên, Vĩnh Phúc cũ), #51 CĐ Cơ điện và Thủy lợi (nghi ở Hưng Yên, cần kiểm tra).",
        "• #42 'CĐ Cơ điện Hà Tây' không xác minh được tên trường. #50 có link website bị đánh dấu nghi vấn.",
        "• Bổ sung: CĐ nghề Việt Nam – Hàn Quốc TP HN (vừa nhận sáp nhập), 2 trường trung cấp đang nâng cấp lên cao đẳng, 14 TT GDNN-GDTX sau sắp xếp, 30 trường đại học.",
    ]),
    ("LƯU Ý THẲNG THẮN", [
        "1. Cột 'Hiện trạng LMS', 'Số hóa quản lý' và khoảng cách trong file gốc không có nguồn, đọc giống nội dung do AI tổng hợp: vài chi tiết sai rõ (vị trí), nhiều chi tiết không kiểm chứng được. Điểm C1 và C10 đang suy ra từ các cột này. Đừng dẫn nguyên văn khi nói chuyện với trường. Gọi điện hoặc xem website để xác minh trước.",
        "2. Điểm của 30 trường đại học là ƯỚC LƯỢNG theo hiểu biết chung (chưa khảo sát LMS). Đại học lớn (HaUI, ĐH Y, ĐH Dược, ĐH Xây dựng) có chu kỳ bán dài và hay đã có hệ thống riêng, mô hình chấm điểm chưa phản ánh hết điều này. Gợi ý: Đợt 1 nên khoảng 2/3 là cao đẳng/trung cấp (thắng nhanh, làm case) và 1/3 là đại học.",
        "3. Bạn chọn 'chào cả gói'. Hợp lý cho bản đề xuất, nhưng ở cuộc gặp đầu nên mở bằng MỘT vấn đề cụ thể của trường (ví dụ hợp nhất học liệu sau sáp nhập, thiếu thiết bị thực hành) và một pilot nhỏ trong 1 học kỳ. Trường công lập khó ký gói lớn khi chưa có trong kế hoạch và dự toán, nên pilot là đường vào kế hoạch 2027.",
        "4. Với trường đã có LMS mạnh (CĐ Công nghệ cao, HCEM, FPT, HaUI…), cột Gói chào tự chuyển sang 'số hóa học liệu chạy trên LMS sẵn có'. Điều này đòi hỏi học liệu/khóa mô phỏng của bạn xuất được chuẩn SCORM/xAPI hoặc LTI. Nếu chưa có, đây là việc kỹ thuật nên làm trước Đợt 2.",
        "5. 14 TT GDNN-GDTX cùng chủ quản là Sở GD&ĐT Hà Nội. Nên đề xuất ở cấp Sở (một nền tảng cho cả cụm) song song với pilot ở 1 trung tâm, thay vì đi lẻ 14 nơi.",
    ]),
]
r = 4
for head, lines in blocks:
    ws0.cell(row=r, column=2, value=head).font = f(True, "FFFFFF")
    ws0.cell(row=r, column=2).fill = HEAD
    r += 1
    for line in lines:
        c = ws0.cell(row=r, column=2, value=line)
        c.font = f(color="555555", size=9, italic=True) if line.strip().startswith("Nguồn") else f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1

# ======================= 2. TRỌNG SỐ =======================
ws1 = wb.create_sheet("Trọng số")
title(ws1, "Khung tiêu chí & trọng số", "Sửa ô vàng. Mọi điểm trong sheet 'Chấm điểm' và phân đợt cập nhật theo.")
header(ws1, 4, ["Mã", "Tiêu chí", "Trọng số (%)", "Đo cái gì", "Điểm 1", "Điểm 3", "Điểm 5"],
       [6, 26, 12, 60, 34, 34, 40])
for i, (code, name, w, what, d1, d3, d5) in enumerate(CRIT):
    rr = 5 + i
    vals = [code, name, w, what, d1, d3, d5]
    for j, v in enumerate(vals, 1):
        c = ws1.cell(row=rr, column=j, value=v)
        c.font = f(bold=(j <= 2))
        c.alignment = WRAP
        c.border = BOX
    ws1.cell(row=rr, column=3).fill = KEYIN if code == "C8" else INPUT
    ws1.cell(row=rr, column=3).font = f(True, "0000FF")
    ws1.cell(row=rr, column=3).alignment = CENTER
ws1.cell(row=16, column=2, value="Tổng trọng số").font = f(True)
ws1["C16"] = "=SUM(C5:C15)"
ws1["C16"].font = f(True)
ws1["C16"].fill = CALC
ws1["C16"].alignment = CENTER
ws1["D16"] = '=IF(C16=100,"OK","Lưu ý: tổng khác 100, điểm vẫn quy về thang 100 nhưng nên giữ 100 cho dễ đọc")'
ws1["D16"].font = f(italic=True, color="555555")
ws1.cell(row=17, column=2, value="Nhóm: Nhu cầu (C1–C4) = 40%; Khả năng mua (C5–C7) = 28%; Tiếp cận (C8–C9) = 18%; Khả thi & lan tỏa (C10–C11) = 14%").font = f(italic=True, color="555555")
ws1.cell(row=19, column=2, value="Ngưỡng Đợt 1 (điểm ≥)").font = f(True)
ws1["C19"] = 70
ws1.cell(row=20, column=2, value="Ngưỡng Đợt 2 (điểm ≥)").font = f(True)
ws1["C20"] = 65
for a in ("C19", "C20"):
    ws1[a].fill = INPUT
    ws1[a].font = f(True, "0000FF")
    ws1[a].alignment = CENTER
ws1["D19"] = "Chọn 70 để Đợt 1 có khoảng 16 trường, vừa sức tiếp cận trong 7 tuần (xem sheet Kế hoạch). Dưới ngưỡng Đợt 2 là Đợt 3. Trường ngoài Hà Nội luôn vào Đợt 4."
ws1["D19"].alignment = WRAP
ws1["D19"].font = f(italic=True, color="555555")
ws1.merge_cells("D19:G20")

# ======================= 3. CHẤM ĐIỂM =======================
ws = wb.create_sheet("Chấm điểm")
crit_cols = list(range(11, 22))  # K..U
labels = ["Hạng", "Trường", "Loại hình", "Chủ quản", "Vùng", "Km từ Hà Đông (file gốc)", "Ngành chính",
          "Hiện trạng LMS (file gốc)", "Số hóa quản lý (file gốc)", "Sự kiện kích hoạt"] + \
         [f"{c[0]}\n{c[1]}" for c in CRIT] + \
         ["Tổng điểm (/100)", "Đợt đề xuất", "Gói chào mở cửa (tự động)", "Độ tin cậy dữ liệu",
          "Căn cứ chấm sơ bộ (từng tiêu chí)", "Ghi chú", "Website", "Nguồn dòng"]
widths = [6, 34, 11, 20, 8, 9, 30, 32, 30, 32] + [9] * 11 + [10, 12, 34, 32, 70, 34, 22, 14]
title(ws, "Chấm điểm & xếp hạng", "Ô vàng là dữ liệu nhập. C8 (cột R) mặc định 1, hãy điền. C9 tự tính từ cột Vùng. Dùng bộ lọc ở hàng tiêu đề để sắp xếp/lọc theo Đợt, Loại hình.")
ws["K3"] = "Trọng số →"
ws["J3"] = "Trọng số →"
ws["J3"].font = f(True, TEAL)
ws["J3"].alignment = Alignment(horizontal="right")
ws["K3"] = None
for k, col in enumerate(crit_cols):
    c = ws.cell(row=3, column=col, value=f"='Trọng số'!C{5 + k}")
    c.font = f(True, "008000")
    c.fill = SUB
    c.alignment = CENTER
header(ws, 4, labels, widths)
ws.row_dimensions[4].height = 58

HN = {  # principal names from Thông báo 750-TB/TU
    "CĐ Y tế Hà Nội": "Hiệu trưởng: Phạm Văn Tân (TB 750-TB/TU)",
    "CĐ nghề Công nghiệp Hà Nội": "Hiệu trưởng: Phạm Thị Hường (TB 750-TB/TU)",
    "CĐ Điện tử - Điện lạnh Hà Nội": "Hiệu trưởng: Trần Xuân Ngọc (TB 750-TB/TU)",
    "CĐ Thương mại và Du lịch Hà Nội": "Hiệu trưởng: Trịnh Thị Thu Hà (TB 750-TB/TU)",
    "CĐ nghề Việt Nam – Hàn Quốc": "Hiệu trưởng: Nguyễn Công Truyền (TB 750-TB/TU)",
    "CĐ Công nghệ cao Hà Nội": "Hiệu trưởng: Phạm Xuân Khánh (TB 750-TB/TU)",
    "CĐ Cộng đồng Hà Tây": "Hiệu trưởng: Hoàng Đức Huế (TB 750-TB/TU)",
    "CĐ Nghệ thuật Hà Nội": "Hiệu trưởng: Vũ Tiến Dũng (TB 750-TB/TU)",
    "Trường ĐH Thủ đô Hà Nội": "Hiệu trưởng: Đỗ Hồng Cường (TB 750-TB/TU)",
}

for i, rw in enumerate(rows):
    R = FIRST + i
    note = rw['note']
    for key, val in HN.items():
        if rw['name'].startswith(key):
            note = (val + ("; " + note if note else ""))
    why = " | ".join(f"{c[0]}={s}: {w}" for c, s, w in zip(CRIT, rw['s'], rw['why']))
    vals = [f"=RANK(V{R},$V${FIRST}:$V${LAST})", rw['name'], rw['loai'], rw['owner'], rw['vung'], rw['km'],
            rw['nganh'], rw['lms'], rw['sis'], rw['trig']]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=R, column=j, value=v)
        c.font = f()
        c.alignment = WRAP
        c.border = BOX
    ws.cell(row=R, column=1).fill = CALC
    ws.cell(row=R, column=1).alignment = CENTER
    ws.cell(row=R, column=2).font = f(True)
    ws.cell(row=R, column=5).fill = INPUT
    ws.cell(row=R, column=5).font = f(color="0000FF")
    ws.cell(row=R, column=5).alignment = CENTER
    for k, col in enumerate(crit_cols):
        if k == 8:  # C9 formula from Vùng
            v = f'=IF(E{R}="Ngoài HN",1,IF(E{R}="","",6-VALUE(E{R})))'
            c = ws.cell(row=R, column=col, value=v)
            c.fill = CALC
            c.font = f()
        else:
            c = ws.cell(row=R, column=col, value=rw['s'][k])
            c.fill = KEYIN if k == 7 else INPUT
            c.font = f(k == 7, "0000FF")
        c.alignment = CENTER
        c.border = BOX
    ws.cell(row=R, column=22, value=f"=ROUND(SUMPRODUCT(K{R}:U{R},$K$3:$U$3)/(5*SUM($K$3:$U$3))*100,1)")
    ws.cell(row=R, column=23, value=f'=IF(E{R}="Ngoài HN","Đợt 4 – Ngoài HN",IF(V{R}>=\'Trọng số\'!$C$19,"Đợt 1",IF(V{R}>=\'Trọng số\'!$C$20,"Đợt 2","Đợt 3")))')
    ws.cell(row=R, column=24, value=f'=IF(K{R}>=4,"LMS + số hóa học liệu",IF(K{R}=3,"Nâng cấp/thay LMS + số hóa học liệu","Số hóa học liệu, chạy trên LMS sẵn có"))&IF(L{R}>=4," + mô phỏng","")&IF(M{R}>=4," + ngoại ngữ","")')
    for col in (22, 23, 24):
        c = ws.cell(row=R, column=col)
        c.fill = CALC
        c.font = f(col == 22)
        c.alignment = CENTER if col != 24 else WRAP
        c.border = BOX
    tail = [rw['trust'], why, note, rw['web'], rw['src']]
    for j, v in enumerate(tail, 25):
        c = ws.cell(row=R, column=j, value=v)
        c.font = f(size=9) if j == 26 else f()
        c.alignment = WRAP
        c.border = BOX
    ws.row_dimensions[R].height = 75

ws.freeze_panes = "C5"
ws.auto_filter.ref = f"A4:{L(len(labels))}{LAST}"
# validations
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                    error="Nhập số nguyên từ 1 đến 5", errorTitle="Điểm 1–5")
ws.add_data_validation(dv)
dv.add(f"K{FIRST}:R{LAST}")
dv.add(f"T{FIRST}:U{LAST}")
dvv = DataValidation(type="list", formula1='"1,2,3,4,Ngoài HN"', allow_blank=True)
ws.add_data_validation(dvv)
dvv.add(f"E{FIRST}:E{LAST}")
# conditional formatting
ws.conditional_formatting.add(f"V{FIRST}:V{LAST}",
                              ColorScaleRule(start_type="min", start_color="F8CBAD", mid_type="percentile", mid_value=50,
                                             mid_color="FFF2CC", end_type="max", end_color="A9D08E"))
for txt, color in [("Đợt 1", "A9D08E"), ("Đợt 2", "DDEBF7"), ("Đợt 3", "EDEDED"), ("Đợt 4", "F4E1F7")]:
    ws.conditional_formatting.add(f"W{FIRST}:W{LAST}",
                                  FormulaRule(formula=[f'LEFT(W{FIRST},5)="{txt}"'], fill=PatternFill("solid", fgColor=color)))
ws.conditional_formatting.add(f"Y{FIRST}:Y{LAST}", FormulaRule(
    formula=[f'OR(ISNUMBER(SEARCH("SAI",Y{FIRST})),ISNUMBER(SEARCH("KHÔNG",Y{FIRST})),ISNUMBER(SEARCH("NGHI",Y{FIRST})))'],
    fill=PatternFill("solid", fgColor="F8CBAD"), stopIfTrue=True))
ws.conditional_formatting.add(f"Y{FIRST}:Y{LAST}", FormulaRule(
    formula=[f'LEFT(Y{FIRST},12)="Đã xác minh"'], fill=PatternFill("solid", fgColor="C6E0B4"), stopIfTrue=True))
ws.conditional_formatting.add(f"Y{FIRST}:Y{LAST}", FormulaRule(
    formula=[f'OR(ISNUMBER(SEARCH("Chưa",Y{FIRST})),ISNUMBER(SEARCH("Ước lượng",Y{FIRST})))'],
    fill=PatternFill("solid", fgColor="FCE4D6")))
ws["R4"].comment = Comment("Điền mức quan hệ của bạn với trường: 1 = chưa quen ai, 3 = quen gián tiếp, 5 = quen trực tiếp BGH/Trưởng phòng Đào tạo.", "Bitsness")

# ======================= 4. PIPELINE =======================
wp = wb.create_sheet("Pipeline")
title(wp, "Pipeline liên hệ", "Đã điền sẵn trường Đợt 1 và Đợt 2 (theo điểm sơ bộ). Thêm dòng mới ở cuối khi cần. Cột Đợt tự cập nhật theo sheet Chấm điểm.")
STAGES = ["1. Chưa liên hệ", "2. Đã liên hệ", "3. Đã gặp", "4. Đã khảo sát nhu cầu", "5. Demo/Pilot",
          "6. Đã gửi báo giá", "7. Ký hợp đồng", "8. Dừng/Mất"]
plabels = ["Trường", "Đợt (tự động)", "Điểm (tự động)", "Người liên hệ", "Chức vụ", "Điện thoại / Email", "Kênh tiếp cận",
           "Ngày liên hệ đầu", "Giai đoạn", "Nhu cầu chính phát hiện", "Gói đề xuất", "Giá trị ước tính (triệu đ)",
           "Bước tiếp theo", "Hạn", "Bước số (tự động)", "Ghi chú"]
header(wp, 4, plabels, [36, 12, 9, 22, 20, 24, 18, 13, 22, 34, 30, 14, 34, 12, 9, 30])
wp.row_dimensions[4].height = 40
# example row
ex = ["VÍ DỤ – CĐ ABC (xóa dòng này khi dùng)", "Đợt 1", 72.0, "Nguyễn Văn A", "Trưởng phòng Đào tạo", "09xx xxx xxx",
      "Giới thiệu", "2026-11-04", "4. Đã khảo sát nhu cầu", "Hợp nhất 2 bộ giáo trình sau sáp nhập; thiếu thiết bị điện",
      "Pilot số hóa 1 khoa + mô phỏng mạch điện", 150, "Gửi đề xuất pilot", "2026-11-15", 4, ""]
for j, v in enumerate(ex, 1):
    c = wp.cell(row=5, column=j, value=v)
    c.font = f(italic=True, color="808080")
    c.border = BOX
pl = [rw for rw in rows if rw['vung'] != "Ngoài HN" and rw['score'] >= 65]
PFIRST = 6
PLAST = PFIRST + len(pl) + 40 - 1
for i, rw in enumerate(pl):
    R = PFIRST + i
    wp.cell(row=R, column=1, value=rw['name']).font = f(True)
    wp.cell(row=R, column=9, value=STAGES[0])
for R in range(PFIRST, PLAST + 1):
    wp.cell(row=R, column=2, value=f"=IFERROR(INDEX('Chấm điểm'!$W${FIRST}:$W${LAST},MATCH(A{R},'Chấm điểm'!$B${FIRST}:$B${LAST},0)),\"\")")
    wp.cell(row=R, column=3, value=f"=IFERROR(INDEX('Chấm điểm'!$V${FIRST}:$V${LAST},MATCH(A{R},'Chấm điểm'!$B${FIRST}:$B${LAST},0)),\"\")")
    wp.cell(row=R, column=15, value=f'=IF(I{R}="","",VALUE(LEFT(I{R},1)))')
    for j in range(1, 17):
        c = wp.cell(row=R, column=j)
        c.border = BOX
        c.alignment = WRAP
        if c.font is None or not c.font.bold:
            c.font = f()
        if j in (2, 3, 15):
            c.fill = CALC
        elif j >= 4 or j == 1:
            c.fill = INPUT
            if j != 1:
                c.font = f(color="0000FF")
    wp.cell(row=R, column=1).font = f(True, "0000FF")
    wp.cell(row=R, column=8).number_format = "dd/mm/yyyy"
    wp.cell(row=R, column=14).number_format = "dd/mm/yyyy"
dvs = DataValidation(type="list", formula1='"' + ",".join(STAGES) + '"', allow_blank=True)
wp.add_data_validation(dvs)
dvs.add(f"I{PFIRST}:I{PLAST}")
dvk = DataValidation(type="list", formula1='"Giới thiệu,Gọi điện,Email,Gặp trực tiếp,Hội thảo/sự kiện,Qua Sở/Bộ chủ quản,Zalo/LinkedIn"', allow_blank=True)
wp.add_data_validation(dvk)
dvk.add(f"G{PFIRST}:G{PLAST}")
wp.freeze_panes = "B5"
wp.auto_filter.ref = f"A4:P{PLAST}"
for txt, color in [("Đợt 1", "A9D08E"), ("Đợt 2", "DDEBF7"), ("Đợt 3", "EDEDED")]:
    wp.conditional_formatting.add(f"B{PFIRST}:B{PLAST}",
                                  FormulaRule(formula=[f'LEFT(B{PFIRST},5)="{txt}"'], fill=PatternFill("solid", fgColor=color)))

# ======================= 5. KẾ HOẠCH =======================
wk = wb.create_sheet("Kế hoạch", 1)
title(wk, "Kế hoạch tiếp cận: Hà Nội trước, ngoài Hà Nội sau", "Mốc thời gian tính từ tuần 19/10/2026. Số trường mỗi đợt và tiến độ phễu tự đếm.")
header(wk, 4, ["Đợt", "Số trường (tự đếm)", "Thời gian", "Mục tiêu đầu ra", "Trọng tâm & cách làm"], [22, 12, 24, 44, 80])
waves = [
    ("Đợt 1", "02/11 – 20/12/2026 (7 tuần)",
     "Gặp ≥10 trường; khảo sát nhu cầu ≥6; gửi ≥3 đề xuất pilot để vào kế hoạch 2027",
     "Ưu tiên trường công lập vừa sáp nhập (BGH mới, phải hợp nhất học liệu) và trường tư thục y–kỹ thuật quyết nhanh. Danh sách tự động hiện là 8 CĐ + 8 ĐH; nên chuyển 2–3 ĐH lớn, chu kỳ dài (HaUI, ĐH Dược, ĐH Xây dựng) sang Đợt 2 và đưa 2 trường trung cấp đang nâng cấp lên Đợt 1. Tận dụng 20/11 (Ngày Nhà giáo VN) để thăm hỏi, kết nối."),
    ("Đợt 2", "07/12/2026 – 24/01/2027",
     "Liên hệ toàn bộ; gặp ≥12; chốt 1–2 pilot chạy học kỳ 2",
     "Chạy song song phần cuối Đợt 1. Dùng phản hồi Đợt 1 chỉnh bộ chào hàng theo phân khúc. Gửi đề xuất cấp cụm cho Sở GD&ĐT về 14 TT GDNN-GDTX."),
    ("Đợt 3", "22/02 – 30/04/2027",
     "Liên hệ theo nhóm qua hội thảo; chuyển pilot Đợt 1 thành hợp đồng",
     "Tổ chức 1 hội thảo chuyên đề (ví dụ 'Hợp nhất học liệu sau sáp nhập' hoặc 'Thực hành ảo cho trường nghề'), mời nhiều trường Đợt 3 cùng lúc, dùng case pilot Đợt 1 làm minh chứng."),
    ("Đợt 4 – Ngoài HN", "05 – 07/2027",
     "Mở rộng sang tỉnh lân cận với case study Hà Nội",
     "Giữ nguyên khung chấm điểm; thay C9 bằng khoảng cách/thời gian di chuyển từ Hà Nội. Ưu tiên tỉnh có cụm trường nghề (Bắc Ninh, Hưng Yên, Phú Thọ, Ninh Bình)."),
]
for i, (w, t, goal, how) in enumerate(waves):
    R = 5 + i
    vals = [w, f"=COUNTIF('Chấm điểm'!$W${FIRST}:$W${LAST},\"{w[:5]}*\")", t, goal, how]
    for j, v in enumerate(vals, 1):
        c = wk.cell(row=R, column=j, value=v)
        c.font = f(j == 1)
        c.alignment = WRAP if j != 2 else CENTER
        c.border = BOX
    wk.cell(row=R, column=2).fill = CALC
    wk.row_dimensions[R].height = 62

wk.cell(row=10, column=1, value="LỊCH CÔNG VIỆC").font = f(True, TEAL, 12)
header(wk, 11, ["Giai đoạn", "Từ", "Đến", "Việc chính", "Đầu ra"])
plan = [
    ("0. Chuẩn bị", "19/10/2026", "01/11/2026",
     "• Xác minh 16 trường Đợt 1: website, gọi tổng đài/phòng Đào tạo, hỏi người quen. Sửa điểm C1, C5, C10.\n• Điền C8 (quan hệ) cho toàn bộ danh sách.\n• Bộ chào hàng theo 7 phân khúc (sheet Kịch bản): 1 trang giới thiệu + 1 đề xuất pilot mẫu.\n• Demo 3 mô-đun: 1 mô phỏng điện/ô tô, 1 mô phỏng điều dưỡng, 1 khóa tiếng Anh chuyên ngành.\n• Hồ sơ năng lực pháp nhân, bảng giá khung, gói pilot 1 học kỳ.\n• Kiểm tra học liệu có xuất SCORM/xAPI/LTI để chạy trên LMS sẵn có của trường.",
     "Danh sách Đợt 1 đã xác minh; bộ tài liệu; demo chạy được"),
    ("1. Đợt 1", "02/11/2026", "20/12/2026",
     "• Mỗi tuần liên hệ 3–4 trường, gặp 2.\n• Cuộc gặp đầu: hỏi theo sheet Câu hỏi khảo sát, không trình bày cả gói ngay.\n• Trong 5 ngày sau gặp: gửi tóm tắt nhu cầu + đề xuất pilot.\n• Cập nhật Pipeline sau mỗi lần liên hệ.",
     "≥10 cuộc gặp, ≥3 đề xuất pilot"),
    ("2. Đợt 2", "07/12/2026", "24/01/2027",
     "• Tương tự Đợt 1, dùng bộ chào hàng đã chỉnh.\n• Gặp Sở GD&ĐT Hà Nội (phòng phụ trách GDNN-GDTX) về đề xuất cấp cụm cho 14 TT.",
     "≥12 cuộc gặp, 1–2 pilot ký"),
    ("Tết Nguyên đán", "25/01/2027", "21/02/2027",
     "• Không mở liên hệ mới (Tết 06/02/2027).\n• Triển khai pilot, số hóa học liệu mẫu, viết case study.",
     "Case study đầu tiên"),
    ("3. Đợt 3 + chuyển đổi", "22/02/2027", "30/04/2027",
     "• Hội thảo chuyên đề cho trường Đợt 3.\n• Đo kết quả pilot, đề xuất hợp đồng năm.",
     "Hợp đồng đầu tiên từ pilot"),
    ("4. Ngoài Hà Nội", "01/05/2027", "31/07/2027",
     "• Chấm lại các trường tỉnh bằng cùng khung.\n• Tiếp cận qua hội thảo/đối tác địa phương.",
     "Danh sách tỉnh + 5–8 cuộc gặp"),
]
for i, row in enumerate(plan):
    R = 12 + i
    for j, v in enumerate(row, 1):
        c = wk.cell(row=R, column=j, value=v)
        c.font = f(j == 1)
        c.alignment = WRAP
        c.border = BOX
    wk.row_dimensions[R].height = 30 + 15 * row[3].count("\n")

wk.cell(row=19, column=1, value="PHỄU TIẾN ĐỘ (đếm từ sheet Pipeline)").font = f(True, TEAL, 12)
header(wk, 20, ["Giai đoạn", "Đã đạt (tự đếm)", "Mục tiêu hết Đợt 2", "Tỷ lệ đạt", "Cách đếm"])
targets = [None, 40, 22, 14, 5, 4, 2]
for k in range(1, 8):
    R = 20 + k
    wk.cell(row=R, column=1, value=STAGES[k - 1] if k > 1 else "1. Trong danh sách")
    if k == 1:
        wk.cell(row=R, column=2, value=f"=COUNTA(Pipeline!$A${PFIRST}:$A${PLAST})")
        wk.cell(row=R, column=5, value="Số trường có trong Pipeline")
    elif k == 7:
        wk.cell(row=R, column=2, value=f"=COUNTIF(Pipeline!$O${PFIRST}:$O${PLAST},7)")
        wk.cell(row=R, column=5, value="Số trường ở giai đoạn 7")
    else:
        wk.cell(row=R, column=2, value=f'=COUNTIFS(Pipeline!$O${PFIRST}:$O${PLAST},">="&{k},Pipeline!$O${PFIRST}:$O${PLAST},"<8")')
        wk.cell(row=R, column=5, value=f"Số trường đã tới giai đoạn {k} trở lên (không tính Dừng/Mất)")
    if targets[k - 1]:
        wk.cell(row=R, column=3, value=targets[k - 1])
        wk.cell(row=R, column=4, value=f"=IF(C{R}=0,\"\",B{R}/C{R})")
        wk.cell(row=R, column=4).number_format = "0%"
    for j in range(1, 6):
        c = wk.cell(row=R, column=j)
        c.border = BOX
        c.font = f(j == 1)
        c.alignment = WRAP if j in (1, 5) else CENTER
    wk.cell(row=R, column=2).fill = CALC
    wk.cell(row=R, column=4).fill = CALC
    wk.cell(row=R, column=3).fill = INPUT
    wk.cell(row=R, column=3).font = f(color="0000FF")
wk.cell(row=28, column=1, value="Mục tiêu phễu là giả định để theo dõi (bạn sửa ô vàng): khoảng 55% trường liên hệ được sẽ gặp, khoảng 60% số đã gặp đi tới khảo sát nhu cầu. Điều chỉnh sau 4 tuần đầu theo số liệu thật.").font = f(italic=True, color="555555")

# ======================= 6. KỊCH BẢN =======================
wsc = wb.create_sheet("Kịch bản")
title(wsc, "Kịch bản tiếp cận theo phân khúc", "Mở cuộc gặp bằng nỗi đau cụ thể, kết thúc bằng một pilot nhỏ. Cả gói nằm trong đề xuất, không trình bày hết ở lần gặp đầu.")
header(wsc, 4, ["Phân khúc", "Trường tiêu biểu trong danh sách", "Nỗi đau thường gặp", "Câu mở đầu gợi ý",
                "Gói mở cửa (pilot)", "Cần gặp ai", "Mang theo bằng chứng gì"], [22, 34, 40, 44, 36, 28, 34])
seg = [
    ("Công lập vừa sáp nhập/nâng cấp", "CĐ Y tế HN; CĐ nghề Công nghiệp HN; CĐ Điện tử-Điện lạnh HN; CĐ Thương mại & Du lịch HN; CĐ nghề Việt Nam–Hàn Quốc; 2 trung cấp nâng cấp",
     "2–3 bộ chương trình, giáo trình, dữ liệu SV khác nhau phải gộp làm một; BGH mới cần kết quả rõ trong năm đầu; giảng viên các cơ sở cũ chưa dùng chung công cụ",
     "\"Sau sáp nhập, trường đang có mấy bộ giáo trình cho cùng một nghề? Chúng tôi giúp kiểm kê và đưa về một kho học liệu chung trên một nền tảng trong một học kỳ.\"",
     "Kiểm kê học liệu miễn phí + pilot số hóa và lên LMS cho 1 khoa/1 nghề trong học kỳ 2",
     "Hiệu trưởng / PHT phụ trách đào tạo; Trưởng phòng Đào tạo; người phụ trách CNTT",
     "Mẫu báo cáo kiểm kê học liệu; demo LMS với lớp gộp từ 2 cơ sở; kế hoạch triển khai 1 học kỳ"),
    ("Y – dược – điều dưỡng", "CĐ Y Dược Tuệ Tĩnh; CĐ Y Dược Pasteur; CĐ Y Dược HN; CĐ Dược HN; CĐ Y HN; CĐ Y tế Bạch Mai; ĐH Đại Nam; ĐH Thăng Long; ĐH Thành Đô",
     "Thiếu giờ thực hành lâm sàng, phòng skills lab đắt; SV điều dưỡng muốn làm việc ở Nhật/Đức cần ngoại ngữ chuyên ngành; tuyển sinh cạnh tranh",
     "\"Sinh viên của trường được luyện bao nhiêu quy trình điều dưỡng trước khi vào viện? Chúng tôi có mô phỏng quy trình để luyện không giới hạn số lần.\"",
     "1 mô-đun mô phỏng điều dưỡng + 1 khóa ngoại ngữ chuyên ngành, chạy thử cho 1 lớp",
     "Trưởng khoa Điều dưỡng/Dược; Phòng Đào tạo; với tư thục: chủ đầu tư/Hiệu trưởng",
     "Video demo mô phỏng; đề cương khóa tiếng Nhật/Đức/Anh điều dưỡng"),
    ("Kỹ thuật: điện – cơ khí – ô tô", "CTECH; HACTECH; CĐ Công nghệ & Thương mại HN; CĐ Bách Khoa; CĐ Đại Việt; CĐ Quốc tế HN; CĐ Kỹ thuật TTB Y tế",
     "Thiết bị thực hành đắt, hao mòn, sĩ số xưởng lớn; giáo trình mô-đun cũ; khó cập nhật công nghệ mới (xe điện, vi mạch)",
     "\"Mỗi học kỳ trường mất bao nhiêu linh kiện, thiết bị vì sinh viên thực hành sai? Mô phỏng giúp SV làm đúng trước khi vào xưởng.\"",
     "Mô phỏng 1 mô-đun (ví dụ mạch điện cơ bản hoặc hệ thống điện ô tô) + số hóa giáo trình chính mô-đun đó",
     "Trưởng khoa Điện/Cơ khí/Ô tô; Phòng Đào tạo",
     "Demo mô phỏng mạch/ô tô; so sánh chi phí vật tư trước và sau"),
    ("Tư thục đa ngành", "CĐ Bách nghệ; CĐ Công nghệ HN (Hitech); HUBT; ĐH Phenikaa; ĐH Hòa Bình",
     "Cạnh tranh tuyển sinh; cần điểm khác biệt để truyền thông; muốn tối ưu chi phí giảng dạy",
     "\"Năm tới trường muốn tuyển sinh bằng điểm khác biệt nào? 'Thực hành ảo + ngoại ngữ ngay từ năm nhất' là thông điệp phụ huynh dễ hiểu.\"",
     "Cả gói, tính phí theo số SV/năm; pilot cho khóa tuyển sinh mới",
     "Chủ đầu tư/Chủ tịch HĐQT; Hiệu trưởng; Trưởng phòng Tuyển sinh",
     "Mẫu nội dung truyền thông tuyển sinh dùng sản phẩm; bảng giá theo SV"),
    ("Đại học", "PTIT; ĐH Xây dựng; ĐH Thủy lợi; UNETI; ĐH Dược HN; ĐH Mở HN; HaUI; ĐH Thủ đô HN",
     "Chuẩn đầu ra tiếng Anh; mục tiêu QĐ 2371 đến 2030; khối lượng giáo trình lớn chưa số hóa; phòng thí nghiệm quá tải",
     "\"Theo QĐ 2371, đến 2030 tối thiểu 20% cơ sở đại học phải đạt mức 1. Trường đang theo lộ trình nào? Chúng tôi có thể số hóa giáo trình và triển khai tiếng Anh chuyên ngành trên nền tảng của trường.\"",
     "Số hóa học liệu 1 khoa + khóa tiếng Anh chuyên ngành; mô phỏng lab cho 1 học phần",
     "Phòng Đào tạo; Trung tâm Ngoại ngữ; Trung tâm Học liệu/Thư viện; Trưởng khoa",
     "Trích QĐ 2371; mẫu học liệu số đã làm; khả năng tích hợp SCORM/LTI với LMS hiện có"),
    ("TT GDNN-GDTX (14 TT sau sắp xếp)", "14 TT khu vực, ưu tiên Hà Đông–Chương Mỹ và Hoàng Mai–Thanh Xuân",
     "Vừa gộp từ 29 TT; học viên GDTX cấp THPT thiếu phòng thí nghiệm Lý–Hóa–Sinh; ngân sách hạn chế; 5 TT sắp chuyển thành trường trung học nghề",
     "\"Học viên GDTX của trung tâm làm thí nghiệm Lý, Hóa, Sinh ở đâu? Phòng thí nghiệm ảo trên một nền tảng dùng chung cho cả 14 trung tâm giúp tiết kiệm đầu tư.\"",
     "Đề xuất cấp cụm tới Sở GD&ĐT; pilot tại 1 TT (Hà Đông–Chương Mỹ)",
     "Sở GD&ĐT HN (phòng phụ trách GDNN-GDTX); Giám đốc TT",
     "Demo lab ảo Lý–Hóa–Sinh; mô hình chi phí cho 14 TT so với đầu tư phòng thí nghiệm thật"),
    ("Đã số hóa mạnh (không chào LMS)", "CĐ Công nghệ cao HN; CĐ Cơ điện HN (HCEM); FPT Polytechnic; ĐH FPT; HaUI",
     "Đã có LMS/Smart Campus; cần nội dung chất lượng cho ngành mới (vi mạch, xe điện, AI)",
     "\"Trường đã có nền tảng tốt. Chúng tôi cung cấp nội dung mô phỏng và học liệu số chạy ngay trên LMS hiện có.\"",
     "Gói nội dung: mô phỏng + học liệu số theo chuẩn SCORM/xAPI cho 1 ngành mới",
     "Trưởng khoa ngành mới; Trung tâm/phòng phụ trách học liệu số",
     "File SCORM chạy thử trên LMS của trường"),
]
for i, row in enumerate(seg):
    R = 5 + i
    for j, v in enumerate(row, 1):
        c = wsc.cell(row=R, column=j, value=v)
        c.font = f(j == 1)
        c.alignment = WRAP
        c.border = BOX
    wsc.row_dimensions[R].height = 120
wsc.freeze_panes = "B5"

# ======================= 7. CÂU HỎI KHẢO SÁT =======================
wq = wb.create_sheet("Câu hỏi khảo sát")
title(wq, "Câu hỏi khảo sát nhu cầu (dùng ở cuộc gặp đầu)", "Mỗi câu gắn với tiêu chí để bạn chấm lại điểm sau buổi gặp.")
header(wq, 4, ["Tiêu chí", "Câu hỏi", "Nghe gì để chấm điểm cao"], [16, 70, 60])
qs = [
    ("C1 LMS", "Hiện trường giao bài, học online, kiểm tra bằng hệ thống gì? Ai quản trị? Hợp đồng với nhà cung cấp đến khi nào?", "Chưa có hệ thống, hoặc dùng Google Classroom/Zalo; hợp đồng cũ sắp hết; không hài lòng"),
    ("C1 / C10", "Trường có tổ CNTT không, mấy người? Giảng viên đã quen dạy online chưa?", "Có 2–3 người CNTT, giảng viên đã dạy online (đủ năng lực triển khai)"),
    ("C2 Mô phỏng", "Mô-đun thực hành nào hay thiếu thiết bị, hỏng thiết bị hoặc tốn vật tư nhất?", "Kể ra được mô-đun cụ thể, có số liệu chi phí vật tư"),
    ("C3 Ngoại ngữ", "Chuẩn đầu ra ngoại ngữ hiện nay là gì? Tỷ lệ SV đạt? Có chương trình đưa SV làm việc ở nước ngoài, hợp tác Hàn/Nhật/Đức không?", "Tỷ lệ đạt thấp; có liên kết nước ngoài; có kế hoạch theo QĐ 2371"),
    ("C4 Học liệu", "Trường có bao nhiêu giáo trình, bài giảng? Đang ở dạng giấy, Word hay PDF? Đã có kế hoạch số hóa chưa?", "Học liệu chủ yếu giấy/PDF; có kế hoạch nhưng chưa có đơn vị làm; vừa gộp nhiều bộ"),
    ("C5 Quy mô", "Trường có bao nhiêu HSSV, bao nhiêu nghề/ngành, bao nhiêu giảng viên?", "Trên 2.000 HSSV, nhiều ngành"),
    ("C6 Tài chính", "Kinh phí cho chuyển đổi số lấy từ đâu (ngân sách, nguồn thu sự nghiệp, chủ đầu tư)? Kế hoạch và dự toán 2027 đã chốt chưa? Trường thường mua theo hình thức nào (mua sắm, thuê dịch vụ CNTT)? Mức nào trường tự quyết được?", "Có nguồn thu tự chủ; dự toán 2027 còn mở; trường quen thuê dịch vụ"),
    ("C7 Kích hoạt", "Năm học tới trường có mục tiêu lớn nào: kiểm định, trường chất lượng cao, mở ngành mới, hoàn thiện sau sáp nhập?", "Có mốc thời gian cụ thể trong 6–12 tháng"),
    ("C8 Tiếp cận", "Ngoài anh/chị, ai sẽ tham gia quyết định? Ai là người quyết cuối cùng?", "Được giới thiệu gặp người quyết định"),
    ("C11 Tham chiếu", "Trường có thuộc mạng lưới, hiệp hội hay cùng chủ quản với trường nào khác? Có sẵn sàng làm case hoặc đồng tổ chức hội thảo không?", "Sẵn sàng chia sẻ kết quả, giới thiệu trường khác"),
]
for i, row in enumerate(qs):
    R = 5 + i
    for j, v in enumerate(row, 1):
        c = wq.cell(row=R, column=j, value=v)
        c.font = f(j == 1)
        c.alignment = WRAP
        c.border = BOX
    wq.row_dimensions[R].height = 45

# sheet order: Đọc trước, Kế hoạch, Trọng số, Chấm điểm, Pipeline, Kịch bản, Câu hỏi
order = ["Đọc trước", "Kế hoạch", "Trọng số", "Chấm điểm", "Pipeline", "Kịch bản", "Câu hỏi khảo sát"]
wb._sheets = [wb[n] for n in order]
for s in wb.worksheets:
    s.sheet_view.zoomScale = 90
wb.active = 3
wb.save(OUT)
print("saved", OUT, N, "rows; pipeline", len(pl))
